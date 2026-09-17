"""Parse and load Scioto County's official sheriff-sale results workbook."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from app.database.session import engine
from pipeline.import_ohio_henry_sales import money


SOURCE_URL = "https://www.sciotocountysheriff.org/sales"
DEFAULT_SOURCE = Path("../tmp/spreadsheets/scioto_oh/2026_real_auction_results.xlsx")
OUTPUT = Path("data/sheriff_sales/oh_scioto_sheriff_sales.json")


def normalize_status(raw: str) -> str:
    value = raw.upper()
    if "CANCEL" in value:
        return "cancelled"
    if "POSTPON" in value or "MOVED" in value:
        return "postponed"
    if "SOLD" in value or "BACK TO PLAINTIFF" in value:
        return "sold"
    if "NO BID" in value:
        return "unsold"
    return "scheduled"


def result_amount(raw: str) -> str | None:
    match = re.search(r"\$\s*([\d,]+(?:\.\d{2})?)", raw)
    return money(match.group(1)) if match else None


def parse_workbook(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    mode = None
    events: list[dict] = []
    current = None
    for row in rows:
        cells = list(row[:7]) + [None] * max(0, 7 - len(row))
        joined = " ".join(str(v) for v in cells if v is not None).upper()
        if "UPCOMING ONLINE SALES" in joined:
            mode = "upcoming"
            continue
        if "SALE RESULTS" in joined and "2026" in joined:
            mode = "results"
            continue
        case = str(cells[0]).strip() if cells[0] else ""
        if re.fullmatch(r"\d{2}-CIE-\d{3}", case, re.I):
            if mode == "upcoming":
                raw_status = str(cells[4] or "UPCOMING")
                date_value = cells[3].date().isoformat() if isinstance(cells[3], datetime) else None
                address, appraisal, attorney = str(cells[5] or "").strip(), money(cells[4]), str(cells[6] or "").strip()
            elif mode == "results":
                raw_status = str(cells[5] or "")
                date_value = cells[3].date().isoformat() if isinstance(cells[3], datetime) else None
                address = str(cells[4] or "").strip() if not isinstance(cells[4], (int, float)) else ""
                appraisal, attorney = None, str(cells[6] or "").strip()
            else:
                continue
            current = {"case": case.upper(), "defendant": str(cells[1] or "").strip(),
                       "parcel": str(cells[2] or "").strip(), "sale_date": date_value,
                       "address": address, "appraisal": appraisal, "attorney": attorney,
                       "raw_status": raw_status, "status": normalize_status(raw_status)}
            events.append(current)
        elif current:
            if cells[1]:
                current["defendant"] += " " + str(cells[1]).strip()
            if cells[2]:
                current["parcel"] += " " + str(cells[2]).strip()
            if cells[4] and not isinstance(cells[4], (int, float)):
                current["address"] += " " + str(cells[4]).strip()

    grouped: dict[str, list[dict]] = {}
    for event in events:
        key = f"{event['case']}|{re.sub(r'\s+', '', event['parcel'])}"
        grouped.setdefault(key, []).append(event)
    records = []
    for key, history in grouped.items():
        dated = [event for event in history if event["sale_date"]]
        latest = max(dated, key=lambda event: event["sale_date"]) if dated else history[-1]
        # Some source rows omit the address; retain it from another event for the same case.
        address = latest["address"] or next((event["address"] for event in reversed(history) if event["address"]), "")
        if not address:
            continue
        sheriff_number = "SCIOTO-" + hashlib.sha256(key.encode()).hexdigest()[:16].upper()
        records.append({
            "state": "OH", "county": "Scioto", "sheriff_number": sheriff_number,
            "court_case_number": latest["case"], "address": " ".join(address.split()),
            "sale_date": latest["sale_date"], "status": latest["status"], "source_url": SOURCE_URL,
            "defendant": latest["defendant"] or None, "plaintiff": None, "upset_price": None,
            "plaintiff_attorney": latest["attorney"] or None,
            "raw_payload": {"parcel_number": latest["parcel"], "appraised_value": latest["appraisal"],
                            "raw_status": latest["raw_status"], "sale_result_amount": result_amount(latest["raw_status"]),
                            "event_history": history, "source_file": path.name},
        })
    return sorted(records, key=lambda record: (record["sale_date"] or "", record["court_case_number"]))


def write_output(records: list[dict], output: Path = OUTPUT) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, indent=2, sort_keys=True) + "\n")


def load(records: list[dict]) -> tuple[int, int]:
    run_id, now = str(uuid.uuid4()), datetime.now(timezone.utc)
    created = updated = 0
    with engine.begin() as connection:
        connection.execute(text("""INSERT INTO scrape_runs(id,job_name,county,source_system,started_at,status,records_found)
          VALUES(:id,'oh_scioto_official_workbook','Scioto','oh_county_official',:now,'running',:count)"""),
          {"id": run_id, "now": now, "count": len(records)})
        for record in records:
            normalized = record["address"]
            city = normalized.rsplit(",", 1)[-1].strip() if "," in normalized else "Unknown"
            street = normalized.rsplit(",", 1)[0].strip()
            address_hash = hashlib.sha256(f"OH|SCIOTO|{normalized.upper()}".encode()).hexdigest()
            property_id = connection.execute(text("""INSERT INTO properties(id,normalized_address,street_address,city,municipality,county,state,address_hash,data_quality_score)
              VALUES(:id,:address,:street,:city,:city,'Scioto','OH',:hash,70)
              ON CONFLICT(address_hash) DO UPDATE SET updated_at=NOW() RETURNING id"""),
              {"id": str(uuid.uuid4()), "address": normalized, "street": street, "city": city, "hash": address_hash}).scalar_one()
            content_hash = hashlib.sha256(json.dumps(record, sort_keys=True).encode()).hexdigest()
            existing = connection.execute(text("SELECT id FROM sheriff_sales WHERE state='OH' AND county='Scioto' AND sheriff_number=:number"), {"number": record["sheriff_number"]}).scalar()
            raw = record["raw_payload"]
            params = {"id": existing or str(uuid.uuid4()), "property_id": property_id, "number": record["sheriff_number"],
                      "case": record["court_case_number"], "defendant": record["defendant"], "sale_date": record["sale_date"],
                      "status": record["status"], "url": record["source_url"], "now": now, "hash": content_hash,
                      "map": raw["parcel_number"], "attorney": record["plaintiff_attorney"], "result": raw["raw_status"],
                      "active": record["status"] in ("scheduled", "postponed")}
            if existing:
                connection.execute(text("""UPDATE sheriff_sales SET property_id=:property_id,court_case_number=:case,defendant=:defendant,current_sale_date=:sale_date,current_status=:status,source_url=:url,last_seen_at=:now,last_scraped_at=:now,raw_source_hash=:hash,map_number=:map,sale_attorney=:attorney,sale_result=:result,is_active=:active,updated_at=NOW() WHERE id=:id"""), params)
                updated += 1
            else:
                connection.execute(text("""INSERT INTO sheriff_sales(id,property_id,state,county,sheriff_number,court_case_number,defendant,current_sale_date,current_status,source_url,source_system,first_seen_at,last_seen_at,last_scraped_at,raw_source_hash,is_active,map_number,sale_attorney,sale_result) VALUES(:id,:property_id,'OH','Scioto',:number,:case,:defendant,:sale_date,:status,:url,'oh_county_official',:now,:now,:now,:hash,:active,:map,:attorney,:result)"""), params)
                created += 1
            for event in raw["event_history"]:
                connection.execute(text("""INSERT INTO sheriff_sale_status_history(id,sheriff_sale_id,status,sale_date,observed_at,source_url,raw_status)
                  SELECT :history_id,:sale_id,:status,:sale_date,:now,:url,:raw_status WHERE NOT EXISTS(
                  SELECT 1 FROM sheriff_sale_status_history WHERE sheriff_sale_id=:sale_id AND status=:status
                  AND sale_date IS NOT DISTINCT FROM :sale_date AND COALESCE(raw_status,'')=COALESCE(:raw_status,''))
                  ON CONFLICT DO NOTHING"""),
                  {"history_id": str(uuid.uuid4()), "sale_id": params["id"], "status": event["status"],
                   "sale_date": event["sale_date"], "now": now, "url": record["source_url"], "raw_status": event["raw_status"]})
            connection.execute(text("""INSERT INTO raw_scrape_records(id,scrape_run_id,state,county,source_record_id,source_url,raw_payload,content_hash,parsing_status,scraped_at) VALUES(:id,:run,'OH','Scioto',:number,:url,CAST(:payload AS JSONB),:hash,'parsed',:now) ON CONFLICT DO NOTHING"""),
              {"id": str(uuid.uuid4()), "run": run_id, "number": record["sheriff_number"], "url": record["source_url"], "payload": json.dumps(record), "hash": content_hash, "now": now})
        connection.execute(text("UPDATE scrape_runs SET completed_at=NOW(),status='completed',records_created=:created,records_updated=:updated WHERE id=:id"), {"created": created, "updated": updated, "id": run_id})
    return created, updated


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--load", action="store_true")
    args = parser.parse_args()
    records = parse_workbook(args.source)
    write_output(records, args.output)
    print(f"parsed {len(records)} unique Scioto County cases")
    if args.load:
        created, updated = load(records)
        print(f"created {created}, updated {updated}")


if __name__ == "__main__":
    main()
