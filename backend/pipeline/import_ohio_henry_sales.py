"""Parse and load Henry County, Ohio's official sheriff-sale archive files."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import text

from app.database.session import engine


ARCHIVE_URL = "https://henrycountyohio.gov/198/Sheriff-Sale-Dates"
SOURCE_DIR = Path("../tmp/pdfs/henry_oh")
OUTPUT = Path("data/sheriff_sales/oh_henry_sheriff_sales.json")
LABELS = ("PLAINTIFF", "DEFENDANT", "CASE #", "ADDRESS", "PARCEL #", "APPRAISED", "START BID", "PLAINTIFF ATTORNEY", "SALE STATUS")


def money(value: str | int | float | None) -> str | None:
    if value in (None, ""):
        return None
    cleaned = re.sub(r"[^0-9.]", "", str(value))
    return str(Decimal(cleaned).quantize(Decimal("0.01"))) if cleaned else None


def normalize_status(raw: str) -> str:
    value = raw.upper()
    if "WITHDRAW" in value or "CANCEL" in value:
        return "cancelled"
    if "POSTPON" in value or "RESCHEDULE" in value:
        return "postponed"
    if "SOLD" in value or "PURCHASED" in value:
        return "sold"
    if "UPCOMING" in value:
        return "scheduled"
    return "unknown"


def sale_date(text_value: str) -> str | None:
    match = re.search(r"TUESDAY,\s+([A-Z]+\s+\d{1,2}[,.]\s*\d{4})", text_value, re.I)
    if not match:
        return None
    raw = match.group(1).replace(".", ",")
    return datetime.strptime(re.sub(r"\s+", " ", raw.title()), "%B %d, %Y").date().isoformat()


def labeled_records(text_value: str) -> list[dict[str, str]]:
    alternatives = "|".join(re.escape(v) for v in sorted(LABELS, key=len, reverse=True))
    marker = re.compile(r"(?m)^\s*(" + alternatives + r"): *(.+?)\s*$")
    matches = list(marker.finditer(text_value))
    records: list[dict[str, str]] = []
    current: dict[str, str] = {}
    for index, match in enumerate(matches):
        label = match.group(1).rstrip(":")
        if label == "PLAINTIFF" and current.get("CASE #"):
            records.append(current)
            current = {}
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text_value)
        continuation = text_value[match.end():end]
        parts = [match.group(2).strip()]
        for line in continuation.splitlines():
            value = line.strip()
            if value and not value.startswith("TUESDAY,") and not value.startswith("ONLINE AT"):
                parts.append(value)
        current[label] = " ".join(filter(None, parts)).strip()
    if current.get("CASE #"):
        records.append(current)
    return records


def parse_pdf(path: Path) -> tuple[str | None, list[dict[str, str]]]:
    text_value = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in PdfReader(path).pages)
    return sale_date(text_value), labeled_records(text_value)


def parse_workbook(path: Path) -> tuple[str | None, list[dict[str, str]]]:
    # One archive link is named “PDF” but serves an XLSX payload.
    book = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    lines = []
    for sheet in book:
        for row in sheet.iter_rows(values_only=True):
            left, right = row[0], row[1] if len(row) > 1 else None
            if left is not None:
                lines.append(f"{left} {right if right is not None else ''}".strip())
    text_value = "\n".join(lines)
    return sale_date(text_value), labeled_records(text_value)


def parse_sources(source_dir: Path = SOURCE_DIR) -> list[dict]:
    snapshots: dict[str, dict] = {}
    paths = sorted(source_dir.glob("*.pdf"))
    for path in paths:
        if path.read_bytes()[:2] == b"PK":
            date_value, rows = parse_workbook(path)
        else:
            date_value, rows = parse_pdf(path)
        for row in rows:
            case = re.sub(r"\s+", "", row["CASE #"]).upper()
            parcel = row.get("PARCEL #", "").strip()
            key = f"HENRY-{case}-{hashlib.sha256(parcel.encode()).hexdigest()[:8].upper()}"
            record = {
                "state": "OH", "county": "Henry", "sheriff_number": key,
                "court_case_number": case, "address": row.get("ADDRESS", "").strip(),
                "sale_date": date_value, "status": normalize_status(row.get("SALE STATUS", "")),
                "upset_price": money(row.get("START BID")), "source_url": ARCHIVE_URL,
                "plaintiff": row.get("PLAINTIFF") or None, "defendant": row.get("DEFENDANT") or None,
                "plaintiff_attorney": row.get("PLAINTIFF ATTORNEY") or None,
                "raw_payload": {"parcel_number": parcel, "appraised_value": money(row.get("APPRAISED")),
                                "starting_bid": money(row.get("START BID")),
                                "raw_status": row.get("SALE STATUS"), "source_file": path.name},
            }
            previous = snapshots.get(key)
            if not previous or (record["sale_date"] or "") >= (previous["sale_date"] or ""):
                snapshots[key] = record
    return sorted(snapshots.values(), key=lambda r: (r["sale_date"] or "", r["sheriff_number"]))


def write_output(records: list[dict], output: Path = OUTPUT) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(records, indent=2, sort_keys=True) + "\n")


def load(records: list[dict]) -> tuple[int, int]:
    run_id, now = str(uuid.uuid4()), datetime.now(timezone.utc)
    created = updated = 0
    with engine.begin() as connection:
        connection.execute(text("""INSERT INTO scrape_runs(id,job_name,county,source_system,started_at,status,records_found)
          VALUES(:id,'oh_henry_official_archive','Henry','oh_county_official',:now,'running',:count)"""),
          {"id": run_id, "now": now, "count": len(records)})
        for record in records:
            normalized = " ".join(record["address"].split())
            city = normalized.rsplit(",", 1)[-1].strip() if "," in normalized else "Unknown"
            street = normalized.rsplit(",", 1)[0].strip()
            address_hash = hashlib.sha256(f"OH|HENRY|{normalized.upper()}".encode()).hexdigest()
            property_id = connection.execute(text("""INSERT INTO properties(id,normalized_address,street_address,city,municipality,county,state,address_hash,data_quality_score)
              VALUES(:id,:address,:street,:city,:city,'Henry','OH',:hash,70)
              ON CONFLICT(address_hash) DO UPDATE SET updated_at=NOW() RETURNING id"""),
              {"id": str(uuid.uuid4()), "address": normalized, "street": street, "city": city, "hash": address_hash}).scalar_one()
            content_hash = hashlib.sha256(json.dumps(record, sort_keys=True).encode()).hexdigest()
            existing = connection.execute(text("SELECT id FROM sheriff_sales WHERE state='OH' AND county='Henry' AND sheriff_number=:number"), {"number": record["sheriff_number"]}).scalar()
            params = {"id": existing or str(uuid.uuid4()), "property_id": property_id, "number": record["sheriff_number"],
                      "case": record["court_case_number"], "plaintiff": record["plaintiff"], "defendant": record["defendant"],
                      "sale_date": record["sale_date"], "status": record["status"], "upset": record["upset_price"],
                      "url": record["source_url"], "now": now, "hash": content_hash,
                      "map": record["raw_payload"]["parcel_number"], "attorney": record["plaintiff_attorney"],
                      "result": record["raw_payload"]["raw_status"]}
            if existing:
                connection.execute(text("""UPDATE sheriff_sales SET property_id=:property_id,court_case_number=:case,plaintiff=:plaintiff,defendant=:defendant,current_sale_date=:sale_date,current_status=:status,upset_price=:upset,source_url=:url,last_seen_at=:now,last_scraped_at=:now,raw_source_hash=:hash,map_number=:map,sale_attorney=:attorney,sale_result=:result,updated_at=NOW() WHERE id=:id"""), params)
                updated += 1
            else:
                connection.execute(text("""INSERT INTO sheriff_sales(id,property_id,state,county,sheriff_number,court_case_number,plaintiff,defendant,current_sale_date,current_status,upset_price,source_url,source_system,first_seen_at,last_seen_at,last_scraped_at,raw_source_hash,is_active,map_number,sale_attorney,sale_result) VALUES(:id,:property_id,'OH','Henry',:number,:case,:plaintiff,:defendant,:sale_date,:status,:upset,:url,'oh_county_official',:now,:now,:now,:hash,TRUE,:map,:attorney,:result)"""), params)
                created += 1
            connection.execute(text("""INSERT INTO raw_scrape_records(id,scrape_run_id,state,county,source_record_id,source_url,raw_payload,content_hash,parsing_status,scraped_at) VALUES(:id,:run,'OH','Henry',:number,:url,CAST(:payload AS JSONB),:hash,'parsed',:now) ON CONFLICT DO NOTHING"""),
              {"id": str(uuid.uuid4()), "run": run_id, "number": record["sheriff_number"], "url": record["source_url"], "payload": json.dumps(record), "hash": content_hash, "now": now})
        connection.execute(text("UPDATE scrape_runs SET completed_at=NOW(),status='completed',records_created=:created,records_updated=:updated WHERE id=:id"), {"created": created, "updated": updated, "id": run_id})
    return created, updated


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=SOURCE_DIR)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--load", action="store_true")
    args = parser.parse_args()
    records = parse_sources(args.source_dir)
    write_output(records, args.output)
    print(f"parsed {len(records)} unique Henry County cases")
    if args.load:
        created, updated = load(records)
        print(f"created {created}, updated {updated}")


if __name__ == "__main__":
    main()
