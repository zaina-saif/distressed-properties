"""Import City of Richmond public assessor and transfer workbooks."""
from __future__ import annotations

import argparse
import hashlib
import tempfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterator

import httpx
from openpyxl import load_workbook
from psycopg2.extras import execute_values
from sqlalchemy import text

from app.database.session import warehouse_engine
from pipeline.public_property_records import decimal_value, integer_value, stable_hash, text_value

SOURCE_ID = "va_richmond_assessor_2026_05"
SOURCE_PAGE = "https://www.rva.gov/assessor-real-estate/data-request"
ASSESSMENT_URL = "https://www.rva.gov/sites/default/files/2026-05/Assessor_Public_Data_Set_2026-05-26.xlsx"
TRANSFERS_URL = "https://www.rva.gov/sites/default/files/2026-05/Assessor_Transfers_2026-05-20.xlsx"
SNAPSHOT_COLUMNS = (
    "source_id", "state", "county", "source_parcel_id", "snapshot_year", "street_address", "city",
    "zip_code", "property_type", "land_use_code", "year_built", "living_area", "land_area",
    "land_area_unit", "bedrooms", "bathrooms", "land_value", "improvement_value",
    "total_assessed_value", "market_value", "source_hash",
)
SNAPSHOT_SQL = f"""
INSERT INTO public_property_snapshots ({','.join(SNAPSHOT_COLUMNS)}) VALUES %s
ON CONFLICT (source_id,source_parcel_id,snapshot_year) DO UPDATE SET
 street_address=EXCLUDED.street_address,city=EXCLUDED.city,zip_code=EXCLUDED.zip_code,
 property_type=EXCLUDED.property_type,land_use_code=EXCLUDED.land_use_code,
 year_built=EXCLUDED.year_built,living_area=EXCLUDED.living_area,
 land_area=EXCLUDED.land_area,land_area_unit=EXCLUDED.land_area_unit,
 bedrooms=EXCLUDED.bedrooms,bathrooms=EXCLUDED.bathrooms,
 land_value=EXCLUDED.land_value,improvement_value=EXCLUDED.improvement_value,
 total_assessed_value=EXCLUDED.total_assessed_value,market_value=EXCLUDED.market_value,
 source_hash=EXCLUDED.source_hash,imported_at=NOW()
WHERE public_property_snapshots.source_hash <> EXCLUDED.source_hash
"""
SALE_COLUMNS = ("source_id", "state", "county", "source_parcel_id", "transaction_id",
                "sale_date", "sale_price", "conveyance_code", "arms_length", "source_hash")
SALE_SQL = f"""
INSERT INTO public_property_sales ({','.join(SALE_COLUMNS)}) VALUES %s
ON CONFLICT (source_id,source_parcel_id,transaction_id) DO UPDATE SET
 sale_date=EXCLUDED.sale_date,sale_price=EXCLUDED.sale_price,
 conveyance_code=EXCLUDED.conveyance_code,arms_length=EXCLUDED.arms_length,
 source_hash=EXCLUDED.source_hash,imported_at=NOW()
WHERE public_property_sales.source_hash <> EXCLUDED.source_hash
"""


def workbook_rows(path: Path) -> Iterator[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [text_value(value) or "" for value in next(rows)]
        for values in rows:
            yield dict(zip(headers, values))
    finally:
        workbook.close()


def parcel_id(value: Any) -> str | None:
    value = text_value(value)
    return f"51760-{value}" if value else None


def valid_year(value: Any, snapshot_year: int) -> int | None:
    year = integer_value(value)
    return year if year and 1600 <= year <= snapshot_year else None


def parse_snapshots(row: dict[str, Any]) -> Iterator[tuple[Any, ...]]:
    identity = parcel_id(row.get("PIN"))
    if not identity:
        return
    address = text_value(row.get("PARCEL_LOCATION"))
    prop_type = text_value(row.get("PROP_TYPE"))
    land_use = text_value(row.get("PRIMARY_USE"))
    half = decimal_value(row.get("HALF_BATH_COUNT"))
    full = decimal_value(row.get("BATH_COUNT"))
    baths = ((full or Decimal(0)) + (half or Decimal(0)) / Decimal(2)
             if full is not None or half is not None else None)
    for position in range(1, 6):
        assessment_date = row.get(f"ASSESSMENT_DATE_{position}")
        if not isinstance(assessment_date, (date, datetime)):
            continue
        snapshot_year = assessment_date.year
        if snapshot_year not in range(2022, 2027):
            continue
        current = snapshot_year == 2026
        total = decimal_value(row.get(f"ASSSESS_TOTAL_VALUE_{position}"))
        core = {
            "source_id": SOURCE_ID, "state": "VA", "county": "Richmond City",
            "source_parcel_id": identity, "snapshot_year": snapshot_year,
            "street_address": address, "city": text_value(row.get("PARCEL_LOCATION_CITY")),
            "zip_code": text_value(row.get("PARCEL_LOCATION_ZIP")), "property_type": prop_type,
            "land_use_code": land_use, "year_built": valid_year(row.get("YEAR_BUILT"), snapshot_year) if current else None,
            "living_area": integer_value(row.get("LIVING_AREA")) if current else None,
            "land_area": decimal_value(row.get("LEGAL_AC")), "land_area_unit": "acres",
            "bedrooms": decimal_value(row.get("BED_COUNT")) if current else None,
            "bathrooms": baths if current else None,
            "land_value": decimal_value(row.get(f"ASSSESS_LAND_VALUE_{position}")),
            "improvement_value": decimal_value(row.get(f"ASSSESS_IMP_VALUE_{position}")),
            "total_assessed_value": total, "market_value": total,
        }
        yield tuple(core[name] for name in SNAPSHOT_COLUMNS[:-1]) + (stable_hash(core),)


def parse_sale(row: dict[str, Any]) -> tuple[Any, ...] | None:
    identity = parcel_id(row.get("PIN"))
    raw_date = row.get("TRANSFER_DATE")
    sale_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date if isinstance(raw_date, date) else None
    price = decimal_value(row.get("CONSIDERATION"))
    if not identity or not sale_date or not date(1800, 1, 1) <= sale_date <= date.today() or price is None or price <= 0:
        return None
    book, page = text_value(row.get("DEED_BOOK")), text_value(row.get("DEED_PAGE"))
    transaction = ":".join(filter(None, (book, page))) or stable_hash({"parcel": identity, "date": sale_date, "price": price})[:32]
    qualified = (text_value(row.get("QUALIFIED")) or "").upper()
    arms_length = True if qualified == "Q" else False if qualified == "U" else None
    conveyance = ";".join(filter(None, (text_value(row.get("SALE_TYPE")),
                                         text_value(row.get("DEED_TYPE")), qualified or None))) or None
    core = {"source_id": SOURCE_ID, "state": "VA", "county": "Richmond City",
            "source_parcel_id": identity, "transaction_id": transaction, "sale_date": sale_date,
            "sale_price": price, "conveyance_code": conveyance, "arms_length": arms_length}
    return tuple(core[name] for name in SALE_COLUMNS[:-1]) + (stable_hash(core),)


def load(path: Path, parser, sql: str, batch_size: int, key) -> int:
    connection = warehouse_engine.raw_connection(); loaded = 0
    try:
        cursor = connection.cursor(); batch = {}
        for source in workbook_rows(path):
            parsed_rows = parser(source)
            if parsed_rows is None:
                continue
            if isinstance(parsed_rows, tuple):
                parsed_rows = (parsed_rows,)
            for parsed in parsed_rows:
                batch[key(parsed)] = parsed
            if len(batch) >= batch_size:
                execute_values(cursor, sql, list(batch.values()), page_size=batch_size); connection.commit()
                loaded += len(batch); batch.clear()
        if batch:
            execute_values(cursor, sql, list(batch.values()), page_size=batch_size); connection.commit(); loaded += len(batch)
    finally:
        connection.close()
    return loaded


def download(url: str, path: Path) -> None:
    if path.exists(): return
    with httpx.stream("GET", url, follow_redirects=True, timeout=180) as response:
        response.raise_for_status()
        with path.open("wb") as output:
            for chunk in response.iter_bytes(1024 * 1024): output.write(chunk)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--assessment", type=Path)
    parser.add_argument("--transfers", type=Path)
    parser.add_argument("--batch-size", type=int, default=5_000)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="richmond_assessor_") as temporary:
        assessment = args.assessment or Path(temporary) / "assessment.xlsx"
        transfers = args.transfers or Path(temporary) / "transfers.xlsx"
        download(ASSESSMENT_URL, assessment); download(TRANSFERS_URL, transfers)
        snapshots = load(assessment, parse_snapshots, SNAPSHOT_SQL, args.batch_size, lambda row: (row[3], row[4]))
        sales = load(transfers, parse_sale, SALE_SQL, args.batch_size, lambda row: (row[3], row[4]))
        with warehouse_engine.begin() as connection:
            for name, path, url, count in (("assessment", assessment, ASSESSMENT_URL, snapshots),
                                           ("transfers", transfers, TRANSFERS_URL, sales)):
                connection.execute(text("""
                    INSERT INTO public_import_files(source_id,file_name,file_url,file_size,sha256,row_count,status,completed_at)
                    VALUES(:source,:name,:url,:size,:sha,:rows,'completed',NOW())
                    ON CONFLICT(source_id,file_name) DO UPDATE SET file_size=EXCLUDED.file_size,
                      sha256=EXCLUDED.sha256,row_count=EXCLUDED.row_count,status='completed',completed_at=NOW()
                """), {"source": SOURCE_ID, "name": name, "url": url, "size": path.stat().st_size,
                       "sha": hashlib.sha256(path.read_bytes()).hexdigest(), "rows": count})
    print(f"complete Richmond snapshots={snapshots:,} sales={sales:,}")


if __name__ == "__main__":
    with warehouse_engine.begin() as connection:
        connection.execute(text("""
            INSERT INTO public_data_sources(source_id,state,county,source_name,source_url,access_method,coverage_notes,last_checked_at)
            VALUES(:id,'VA','Richmond City','City of Richmond Assessor Public Data',:url,'xlsx',
              '2022-2026 assessment history and 2015-current transfers; owner, mailing and party-name fields excluded',NOW())
            ON CONFLICT(source_id) DO UPDATE SET last_checked_at=NOW(),updated_at=NOW()
        """), {"id": SOURCE_ID, "url": SOURCE_PAGE})
    main()
