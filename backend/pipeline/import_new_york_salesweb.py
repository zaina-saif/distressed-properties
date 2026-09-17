"""Import user-downloaded official NY Sales Web CSV/XLSX exports.

The state portal requires an interactive CAPTCHA. This module deliberately starts
from a file exported through that public interface and does not automate CAPTCHA.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from psycopg2.extras import execute_values
from sqlalchemy import text

from app.database.session import warehouse_engine
from pipeline.import_new_york_assessment_rolls import parcel_id
from pipeline.public_property_records import date_value, decimal_value, stable_hash, text_value


SOURCE_ID = "ny_orpts_salesweb_rp5217"
SOURCE_URL = "https://pad.tax.ny.gov/index"
SALE_COLUMNS = ("source_id", "state", "county", "source_parcel_id", "transaction_id", "sale_date",
                "sale_price", "recording_date", "conveyance_code", "arms_length", "source_hash")
SNAPSHOT_COLUMNS = ("source_id", "state", "county", "source_parcel_id", "snapshot_year", "street_address",
                    "city", "zip_code", "property_type", "land_use_code", "school_district_code",
                    "school_district_name", "source_hash")
SALE_SQL = f"""INSERT INTO public_property_sales ({','.join(SALE_COLUMNS)}) VALUES %s
ON CONFLICT(source_id,source_parcel_id,transaction_id) DO UPDATE SET sale_date=EXCLUDED.sale_date,
sale_price=EXCLUDED.sale_price,recording_date=EXCLUDED.recording_date,
conveyance_code=EXCLUDED.conveyance_code,arms_length=EXCLUDED.arms_length,
source_hash=EXCLUDED.source_hash,imported_at=NOW()
WHERE public_property_sales.source_hash<>EXCLUDED.source_hash"""
SNAPSHOT_SQL = f"""INSERT INTO public_property_snapshots ({','.join(SNAPSHOT_COLUMNS)}) VALUES %s
ON CONFLICT(source_id,source_parcel_id,snapshot_year) DO UPDATE SET street_address=EXCLUDED.street_address,
city=EXCLUDED.city,zip_code=EXCLUDED.zip_code,property_type=EXCLUDED.property_type,
land_use_code=EXCLUDED.land_use_code,school_district_code=EXCLUDED.school_district_code,
school_district_name=EXCLUDED.school_district_name,source_hash=EXCLUDED.source_hash,imported_at=NOW()
WHERE public_property_snapshots.source_hash<>EXCLUDED.source_hash"""


def key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def normalized(row: dict[Any, Any]) -> dict[str, Any]:
    return {key(name): value for name, value in row.items() if name is not None}


def pick(row: dict[str, Any], *names: str) -> Any:
    return next((row[name] for name in names if name in row and row[name] not in (None, "")), None)


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = date_value(value)
    if parsed:
        return parsed
    raw = text_value(value)
    if raw:
        for pattern in ("%m/%d/%Y", "%m/%d/%y", "%Y%m%d"):
            try:
                return datetime.strptime(raw, pattern).date()
            except ValueError:
                pass
    return None


def bool_value(value: Any) -> bool | None:
    raw = (text_value(value) or "").lower()
    if raw in {"y", "yes", "true", "1", "a", "arms length", "arm's length"}:
        return True
    if raw in {"n", "no", "false", "0"}:
        return False
    return None


def parse(row_value: dict[Any, Any]) -> tuple[tuple[Any, ...], tuple[Any, ...]] | None:
    row = normalized(row_value)
    county = text_value(pick(row, "county_name", "county_nam", "county"))
    swis = text_value(pick(row, "swis_code", "swis_cd", "swis", "municipality_code"))
    print_key = text_value(pick(row, "print_key", "tax_map_id", "parcel_id", "parcel_number"))
    sold = parse_date(pick(row, "sale_date", "sale_dte", "date_of_sale", "transfer_date"))
    price = decimal_value(pick(row, "sale_price", "full_sale_price", "total_sale_price", "price"))
    parcel = parcel_id(swis, print_key)
    if not county or not parcel or not sold or sold.year < 1900 or price is None or price <= 0:
        return None
    book = text_value(pick(row, "deed_book", "book", "liber"))
    page = text_value(pick(row, "deed_page", "page"))
    document = text_value(pick(row, "document_id", "instrument_number", "recording_number"))
    transaction_id = document or (f"{book}:{page}" if book or page else hashlib.sha256(
        f"{parcel}|{sold}|{price}".encode()).hexdigest()[:24])
    prop_class = text_value(pick(row, "property_class_at_sale", "prop_class_at_sale",
                                 "property_class_on_roll", "prop_class_last_roll"))
    property_type = text_value(pick(row, "prop_class_cd_desc_sale", "prop_class_cd_desc_roll")) or prop_class
    address = text_value(pick(row, "property_address", "street_address", "address"))
    if not address:
        address = " ".join(filter(None, (text_value(row.get("st_nbr")), text_value(row.get("st_nam"))))) or None
    sale_core = {"source_id": SOURCE_ID, "state": "NY", "county": county,
                 "source_parcel_id": parcel, "transaction_id": transaction_id,
                 "sale_date": sold, "sale_price": price,
                 "recording_date": parse_date(pick(row, "recording_date", "deed_date", "deed_dte", "date_recorded")),
                 "conveyance_code": text_value(pick(row, "deed_type", "conveyance_code", "sale_condition")),
                 "arms_length": bool_value(pick(row, "arms_length", "arms_length_flag", "rar_usable"))}
    snapshot_core = {"source_id": SOURCE_ID, "state": "NY", "county": county,
                     "source_parcel_id": parcel, "snapshot_year": sold.year,
                     "street_address": address,
                     "city": text_value(pick(row, "city", "property_city", "municipality_name", "muni_name", "muni_nam")),
                     "zip_code": text_value(pick(row, "property_zip_code", "zip_code", "zip", "zip5")),
                     "property_type": property_type, "land_use_code": prop_class,
                     "school_district_code": text_value(pick(row, "school_district_code", "school_cd")),
                     "school_district_name": text_value(pick(row, "school_district_name", "school_nam"))}
    sale = tuple(sale_core[name] for name in SALE_COLUMNS[:-1]) + (stable_hash(sale_core),)
    snapshot = tuple(snapshot_core[name] for name in SNAPSHOT_COLUMNS[:-1]) + (stable_hash(snapshot_core),)
    return sale, snapshot


def rows(path: Path) -> Iterator[dict[Any, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            yield from csv.DictReader(source)
        return
    book = load_workbook(path, read_only=True, data_only=True)
    for sheet in book:
        values = sheet.iter_rows(values_only=True)
        try:
            headers = next(values)
        except StopIteration:
            continue
        for values_row in values:
            yield dict(zip(headers, values_row))


def import_file(path: Path, dry_run: bool = False, batch_size: int = 10000) -> tuple[int, int, int]:
    sale_batch: dict[tuple[str, str], tuple[Any, ...]] = {}
    snapshot_batch: dict[tuple[str, int], tuple[Any, ...]] = {}
    rejected = sales_total = snapshots_total = 0
    raw = warehouse_engine.raw_connection() if not dry_run else None
    try:
        for row in rows(path):
            record = parse(row)
            if not record:
                rejected += 1
                continue
            sale, snapshot = record
            sale_batch[(sale[3], sale[4])] = sale
            snapshot_batch[(snapshot[3], snapshot[4])] = snapshot
            if raw is not None and len(sale_batch) >= batch_size:
                cursor = raw.cursor()
                execute_values(cursor, SALE_SQL, list(sale_batch.values()), page_size=batch_size)
                execute_values(cursor, SNAPSHOT_SQL, list(snapshot_batch.values()), page_size=batch_size)
                raw.commit()
                sales_total += len(sale_batch); snapshots_total += len(snapshot_batch)
                sale_batch.clear(); snapshot_batch.clear()
        sales_total += len(sale_batch); snapshots_total += len(snapshot_batch)
        if raw is not None and sale_batch:
            cursor = raw.cursor()
            execute_values(cursor, SALE_SQL, list(sale_batch.values()), page_size=batch_size)
            execute_values(cursor, SNAPSHOT_SQL, list(snapshot_batch.values()), page_size=batch_size)
            raw.commit()
        return sales_total, snapshots_total, rejected
    finally:
        if raw is not None:
            raw.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.dry_run:
        with warehouse_engine.begin() as connection:
            connection.execute(text("""INSERT INTO public_data_sources
              (source_id,state,source_name,source_url,access_method,coverage_notes,last_checked_at)
              VALUES(:id,'NY','ORPTS Sales Web RP-5217',:url,'Manual public portal export; local batch import',
              'Ten years of non-NYC transfers; party names excluded; portal CAPTCHA is not automated',NOW())
              ON CONFLICT(source_id) DO UPDATE SET last_checked_at=NOW(),updated_at=NOW()"""),
              {"id": SOURCE_ID, "url": SOURCE_URL})
    for path in args.files:
        sales, snapshots, rejected = import_file(path, args.dry_run)
        print(f"{path}: sales={sales:,} snapshots={snapshots:,} rejected_rows={rejected:,}")


if __name__ == "__main__":
    main()
