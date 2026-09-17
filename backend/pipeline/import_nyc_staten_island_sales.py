"""Import NYC Finance annual/rolling Staten Island sales workbooks into the warehouse."""
from __future__ import annotations

import argparse
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from psycopg2.extras import execute_values
from sqlalchemy import text

from app.database.session import warehouse_engine
from pipeline.public_property_records import decimal_value, integer_value, stable_hash, text_value

SOURCE_ID = "nyc_dof_staten_island_annual_sales"
SOURCE_URL = "https://www.nyc.gov/site/finance/property/property-annualized-sales-update.page"
COUNTY = "Richmond"
SALE_COLUMNS = ("source_id", "state", "county", "source_parcel_id", "transaction_id", "sale_date",
                "sale_price", "recording_date", "conveyance_code", "arms_length", "source_hash")
SNAPSHOT_COLUMNS = ("source_id", "state", "county", "source_parcel_id", "snapshot_year", "street_address",
                    "zip_code", "property_type", "land_use_code", "year_built", "living_area", "land_area",
                    "land_area_unit", "source_hash")
SALE_SQL = f"""INSERT INTO public_property_sales ({','.join(SALE_COLUMNS)}) VALUES %s
ON CONFLICT(source_id,source_parcel_id,transaction_id) DO UPDATE SET
sale_date=EXCLUDED.sale_date,sale_price=EXCLUDED.sale_price,conveyance_code=EXCLUDED.conveyance_code,
source_hash=EXCLUDED.source_hash,imported_at=NOW()
WHERE public_property_sales.source_hash<>EXCLUDED.source_hash"""
SNAPSHOT_SQL = f"""INSERT INTO public_property_snapshots ({','.join(SNAPSHOT_COLUMNS)}) VALUES %s
ON CONFLICT(source_id,source_parcel_id,snapshot_year) DO UPDATE SET
street_address=EXCLUDED.street_address,zip_code=EXCLUDED.zip_code,property_type=EXCLUDED.property_type,
land_use_code=EXCLUDED.land_use_code,year_built=EXCLUDED.year_built,living_area=EXCLUDED.living_area,
land_area=EXCLUDED.land_area,land_area_unit=EXCLUDED.land_area_unit,source_hash=EXCLUDED.source_hash,
imported_at=NOW()
WHERE public_property_snapshots.source_hash<>EXCLUDED.source_hash"""


def normalize_header(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def parse_row(row: dict[str, object]):
    if str(row.get("BOROUGH", "")).strip() not in {"5", "5.0"}:
        return None
    block, lot = integer_value(row.get("BLOCK")), integer_value(row.get("LOT"))
    price = decimal_value(row.get("SALEPRICE"))
    raw_date = row.get("SALEDATE")
    sale_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date if isinstance(raw_date, date) else None
    if not block or not lot or not sale_date or price is None or price <= 0:
        return None
    parcel = f"NYC:BBL:5{block:05d}{lot:04d}"
    address = text_value(row.get("ADDRESS"))
    unit = text_value(row.get("APARTMENTNUMBER"))
    sale_class = text_value(row.get("BUILDINGCLASSATTIMEOFSALE"))
    # The city files have no deed/document ID; use stable transaction details for
    # idempotence across the annual and rolling files, without implying deed identity.
    transaction_id = stable_hash({"parcel": parcel, "date": sale_date, "price": price,
                                  "unit": unit, "address": address})
    sale_core = dict(source_id=SOURCE_ID, state="NY", county=COUNTY, source_parcel_id=parcel,
                     transaction_id=transaction_id, sale_date=sale_date, sale_price=price,
                     recording_date=None, conveyance_code=sale_class, arms_length=None)
    year_built = integer_value(row.get("YEARBUILT"))
    snapshot_core = dict(source_id=SOURCE_ID, state="NY", county=COUNTY, source_parcel_id=parcel,
                         snapshot_year=sale_date.year, street_address=address,
                         zip_code=text_value(row.get("ZIPCODE")),
                         property_type=text_value(row.get("BUILDINGCLASSCATEGORY")),
                         land_use_code=text_value(row.get("BUILDINGCLASSATPRESENT")),
                         year_built=year_built if year_built and 1600 <= year_built <= 2100 else None,
                         living_area=integer_value(row.get("GROSSSQUAREFEET")),
                         land_area=decimal_value(row.get("LANDSQUAREFEET")),
                         land_area_unit="sqft")
    return (tuple(sale_core[c] for c in SALE_COLUMNS[:-1]) + (stable_hash(sale_core),),
            tuple(snapshot_core[c] for c in SNAPSHOT_COLUMNS[:-1]) + (stable_hash(snapshot_core),))


def workbook_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = None
        for cells in sheet.iter_rows(values_only=True):
            if headers is None:
                candidate = [normalize_header(value) for value in cells]
                if "BOROUGH" in candidate and "SALEDATE" in candidate and "SALEPRICE" in candidate:
                    headers = candidate
                continue
            row = dict(zip(headers, cells))
            parsed = parse_row(row)
            if parsed:
                yield parsed
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    for path in args.files:
        if not path.is_file() or path.suffix.lower() != ".xlsx":
            parser.error(f"not an xlsx file: {path}")
    if not args.dry_run:
        with warehouse_engine.begin() as connection:
            connection.execute(text("""INSERT INTO public_data_sources
              (source_id,state,county,source_name,source_url,access_method,coverage_notes,last_checked_at)
              VALUES(:id,'NY','Richmond','NYC DOF Staten Island annual and rolling sales',:url,'Official XLSX download',
              'Positive-price Staten Island sales; no deed ID or arms-length designation; one snapshot per BBL/year',NOW())
              ON CONFLICT(source_id) DO UPDATE SET last_checked_at=NOW(),updated_at=NOW()"""),
                               {"id": SOURCE_ID, "url": SOURCE_URL})
    raw = warehouse_engine.raw_connection() if not args.dry_run else None
    try:
        for path in args.files:
            sales = {}
            snapshots = {}
            for sale, snapshot in workbook_rows(path):
                sales[(sale[3], sale[4])] = sale
                snapshots[(snapshot[3], snapshot[4])] = snapshot
            if raw is not None:
                cursor = raw.cursor()
                try:
                    for rows, sql in ((sales, SALE_SQL), (snapshots, SNAPSHOT_SQL)):
                        values = list(rows.values())
                        for start in range(0, len(values), 5000):
                            execute_values(cursor, sql, values[start:start + 5000], page_size=5000)
                    raw.commit()
                except Exception:
                    raw.rollback()
                    raise
                finally:
                    cursor.close()
            print(f"{path.name}: qualifying_sales={len(sales):,} snapshots={len(snapshots):,}", flush=True)
    finally:
        if raw is not None:
            raw.close()


if __name__ == "__main__":
    main()
